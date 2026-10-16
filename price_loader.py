# @Version  : 1.0
# @Author   : 小野
# @File     : price_loader.py 单价表读取
# @Time     : 2026/9/30 13:18

import os
import re
from openpyxl import load_workbook


SUFFIXES = ["帽檐", "下巴", "尾翼", "顶盖"]


def _norm_key(m):
    s = str(m).strip().replace(" ", "").upper()
    if not s:
        return s
    try:
        f = float(s)
        if f == int(f):
            return str(int(f))
        return str(f)
    except (ValueError, TypeError):
        pass
    return s


class PriceTable:
    def __init__(self, file_path=None):
        self.file_path = file_path or ""
        self.data = {}
        self.loaded = False
        self.error = ""
        self._last_mtime = 0

    def set_path(self, path):
        self.file_path = path or ""
        self._last_mtime = 0

    def load(self):
        self.data = {}
        self.loaded = False
        self.error = ""

        if not self.file_path:
            self.error = "未设置单价表路径"
            return False

        if not os.path.exists(self.file_path):
            self.error = f"单价表文件不存在：{self.file_path}"
            return False

        try:
            wb = load_workbook(self.file_path, data_only=True, read_only=True)
        except Exception as e:
            self.error = f"读取单价表失败：{e}"
            return False

        try:
            if "Sheet1" not in wb.sheetnames:
                self.error = "单价表中没有 Sheet1"
                wb.close()
                return False

            ws = wb["Sheet1"]
            for row in ws.iter_rows(min_row=3, values_only=True):
                if not row:
                    continue
                model_cell = row[0] if len(row) > 0 else None
                base_cell = row[1] if len(row) > 1 else None
                suffix_cell = row[2] if len(row) > 2 else None

                if model_cell is None:
                    continue
                model_str = str(model_cell).strip()
                if model_str == "":
                    continue

                base_price = None
                if base_cell is not None:
                    try:
                        base_price = float(base_cell)
                    except (ValueError, TypeError):
                        base_price = None

                suffix_dict = self._parse_suffix(suffix_cell)
                models = self._split_models(model_str)

                for m in models:
                    m_norm = _norm_key(m)
                    if m_norm == "":
                        continue
                    self.data[m_norm] = {
                        "base": base_price,
                        "suffix": suffix_dict,
                    }

            wb.close()
            self.loaded = True

            try:
                self._last_mtime = os.path.getmtime(self.file_path)
            except Exception:
                self._last_mtime = 0

            print(f"单价表加载成功，共 {len(self.data)} 个型号")
            return True
        except Exception as e:
            self.error = f"解析单价表失败：{e}"
            try:
                wb.close()
            except Exception:
                pass
            return False

    def reload_if_changed(self):
        if not self.file_path or not os.path.exists(self.file_path):
            return False
        try:
            mtime = os.path.getmtime(self.file_path)
        except Exception:
            return False
        if mtime <= self._last_mtime:
            return False
        print("检测到单价表文件已更新，自动重载...")
        return self.load()

    def _split_models(self, s):
        s = s.replace("（", "(").replace("）", ")")
        s = s.replace("，", ",").replace("、", ",")

        result = []
        m = re.match(r"^([^(]+)", s)
        if m:
            part = m.group(1).strip()
            if part:
                result.append(part)

        for inner in re.findall(r"\(([^)]+)\)", s):
            for part in inner.split(","):
                p = part.strip()
                if p:
                    result.append(p)

        return result

    def _parse_suffix(self, cell):
        if cell is None:
            return {}
        s = str(cell)
        s = s.replace("<br>", " ").replace("<br/>", " ").replace("<BR>", " ")
        s = s.replace("，", " ").replace(",", " ").replace("、", " ")
        s = s.replace("；", " ").replace(";", " ")
        s = s.replace("\n", " ").replace("\r", " ")

        result = {}
        pattern = r"(帽檐|下巴|尾翼|顶盖)\s*([\d.]+)"
        for m in re.finditer(pattern, s):
            kw = m.group(1)
            try:
                result[kw] = float(m.group(2))
            except ValueError:
                pass
        return result

    def get_price(self, model):
        if not self.loaded:
            return None
        model = str(model).strip().replace(" ", "")
        if model == "":
            return None

        for suf in SUFFIXES:
            if model.endswith(suf):
                base = model[:-len(suf)]
                base_key = _norm_key(base)
                if base_key in self.data:
                    suffix_dict = self.data[base_key]["suffix"]
                    if suf in suffix_dict:
                        return suffix_dict[suf]
                return None

        key = _norm_key(model)
        if key in self.data:
            return self.data[key]["base"]
        return None